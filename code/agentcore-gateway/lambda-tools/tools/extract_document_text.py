"""Extract plain text from an uploaded requirements document.

Implements the Document_Extractor (design Section 7b, requirements 5.2-5.3).

Behavior:

* ``.md`` / ``.txt`` -> the document body is returned directly (no extraction
  needed). Requirement 5.2 allows direct ingestion for these formats.
* ``.pdf`` -> text is extracted with :mod:`pypdf`.
* ``.docx`` -> text is extracted with :mod:`docx` (``python-docx``).

The normalized plain text is written back to S3 alongside the source file at::

    {project_id}/aidlc-docs/inception/requirements/uploaded/extracted/<filename>.txt

and also returned to the caller so the Requirements_Analysis_Agent can ingest
it without a second round trip.

Inputs (Gateway MCP tool input):
    file_path: The S3 key of the uploaded document. Must live under a
        ``.../uploaded/`` prefix; the extracted output is placed under
        ``.../uploaded/extracted/`` next to it (so a downstream agent can list
        the prefix and find the normalized version).
"""

import io
import os

from botocore.exceptions import ClientError
from shared import BUCKET_NAME, s3_client, success, error


# ---------------------------------------------------------------------------
# Per-format extractors
# ---------------------------------------------------------------------------


def _extract_text_from_pdf(body: bytes) -> str:
    """Extract text from a PDF document using ``pypdf``.

    Concatenates the text of each page with a single blank line separator so
    page boundaries are preserved without leaking PDF-specific artifacts.
    """
    # Imported lazily so unit tests for the .md / .txt path can run without
    # the binary-heavy PDF dependency installed.
    from pypdf import PdfReader  # type: ignore

    reader = PdfReader(io.BytesIO(body))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    return "\n\n".join(pages)


def _extract_text_from_docx(body: bytes) -> str:
    """Extract text from a DOCX document using ``python-docx``.

    Concatenates paragraph text with newline separators. Tables are flattened
    row-by-row with tab-separated cells so downstream agents see a stable,
    parseable plain-text layout.
    """
    # Imported lazily so unit tests for the .md / .txt path can run without
    # the binary-heavy DOCX dependency installed.
    from docx import Document  # type: ignore

    document = Document(io.BytesIO(body))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            parts.append("\t".join(cells))
    return "\n".join(parts)


_EXTRACTORS = {
    ".pdf": _extract_text_from_pdf,
    ".docx": _extract_text_from_docx,
}
_PASSTHROUGH = {".md", ".txt"}


# ---------------------------------------------------------------------------
# Path helpers
# ---------------------------------------------------------------------------


def _extracted_key(file_path: str) -> str:
    """Return the S3 key for the normalized text alongside ``file_path``.

    The handler expects the source key to live directly under an
    ``.../uploaded/`` prefix; that is how the Frontend Project API issues
    presigned PUTs (design Section 7a). The extracted output is placed under
    ``.../uploaded/extracted/<filename>.txt`` so a single ``list_files`` over
    ``.../uploaded/`` discovers both the source and its normalized form.

    Raises ``ValueError`` if ``file_path`` does not contain a filename or its
    immediate parent directory is not named ``uploaded``.
    """
    directory, filename = os.path.split(file_path)
    if not filename:
        raise ValueError("file_path must include a filename")
    parent_name = os.path.basename(directory.rstrip("/"))
    if parent_name != "uploaded":
        raise ValueError(
            "file_path must live directly under an 'uploaded/' directory"
        )
    base, _ext = os.path.splitext(filename)
    return f"{directory}/extracted/{base}.txt"


# ---------------------------------------------------------------------------
# Tool entrypoint
# ---------------------------------------------------------------------------


def extract_document_text(tool_input):
    """Read an uploaded document from S3 and return its plain text.

    For ``.pdf`` and ``.docx`` inputs the extracted text is also written back to
    S3 at ``.../uploaded/extracted/<filename>.txt``. For ``.md`` and ``.txt``
    inputs the body is also written to the ``extracted/`` location so the
    Requirements_Analysis_Agent has one consistent place to read normalized
    text from regardless of the source format.
    """
    file_path = tool_input.get("file_path")
    if not file_path:
        return error("Missing required field: file_path")

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in _PASSTHROUGH and ext not in _EXTRACTORS:
        return error(
            f"Unsupported document format: {ext or '(none)'}. "
            "Supported formats are .md, .txt, .pdf, .docx"
        )

    try:
        try:
            extracted_key = _extracted_key(file_path)
        except ValueError as ve:
            return error(str(ve))

        try:
            response = s3_client.get_object(Bucket=BUCKET_NAME, Key=file_path)
        except ClientError as e:
            if e.response["Error"]["Code"] == "NoSuchKey":
                return error(f"File not found: {file_path}")
            return error(f"Failed to read file: {e}")

        body = response["Body"].read()

        if ext in _PASSTHROUGH:
            try:
                text = body.decode("utf-8")
            except UnicodeDecodeError:
                # Fall back to a lossy decode rather than 500-ing — uploads
                # may originate from editors that emit non-UTF-8 bytes.
                text = body.decode("utf-8", errors="replace")
        else:
            try:
                text = _EXTRACTORS[ext](body)
            except Exception as exc:  # noqa: BLE001 — surface any extractor error
                return error(f"Failed to extract text from {file_path}: {exc}")

        try:
            s3_client.put_object(
                Bucket=BUCKET_NAME,
                Key=extracted_key,
                Body=text.encode("utf-8"),
                ContentType="text/plain; charset=utf-8",
            )
        except ClientError as e:
            return error(f"Failed to write extracted text: {e}")

        return success(text)
    except ClientError as e:
        return error(f"S3 operation failed: {e}")
